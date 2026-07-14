#!/usr/bin/env python3
"""Convert docs/guides/使用说明.md to docs/guides/使用说明.docx"""

import re
import os
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# Paths
SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
MD_PATH = PROJECT_ROOT / "docs" / "guides" / "使用说明.md"
OUTPUT_PATH = PROJECT_ROOT / "docs" / "guides" / "使用说明.docx"
SCREENSHOTS_DIR = PROJECT_ROOT / "docs" / "guides"


def set_cell_shading(cell, color: str):
    """Set cell background color."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def add_horizontal_line(doc):
    """Add a horizontal divider line using bottom border on an empty paragraph."""
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        '  <w:bottom w:val="single" w:sz="6" w:space="1" w:color="auto"/>'
        '</w:pBdr>'
    )
    pPr.append(pBdr)
    return p


def apply_inline_formatting(paragraph, text, base_font_name="Arial", base_font_size=Pt(11)):
    """Parse inline **bold** and `code` formatting and add runs accordingly."""
    # Pattern: split by **bold** and `code` markers
    pattern = re.compile(r'(\*\*.*?\*\*|`[^`]+`)')
    parts = pattern.split(text)

    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = base_font_name
            run.font.size = base_font_size
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            # light gray background via highlight
            from docx.enum.text import WD_COLOR_INDEX
            run.font.highlight_color = WD_COLOR_INDEX.GRAY_25
        else:
            run = paragraph.add_run(part)
            run.font.name = base_font_name
            run.font.size = base_font_size


def parse_table(lines):
    """Parse markdown table lines into rows of cells. Returns list of lists."""
    rows = []
    for line in lines:
        line = line.strip()
        if re.match(r'^\|[\s\-:]+\|', line):
            continue  # separator row
        cells = [c.strip() for c in line.strip('|').split('|')]
        rows.append(cells)
    return rows


def add_table_to_doc(doc, rows):
    """Add a table to the document."""
    if not rows:
        return
    num_cols = len(rows[0])
    table = doc.add_table(rows=len(rows), cols=num_cols)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, row_data in enumerate(rows):
        for j, cell_text in enumerate(row_data):
            cell = table.rows[i].cells[j]
            # Clear default paragraph
            cell.paragraphs[0].clear()
            paragraph = cell.paragraphs[0]
            # Apply inline formatting
            apply_inline_formatting(paragraph, cell_text)

            if i == 0:
                # Header row styling
                set_cell_shading(cell, "4472C4")
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    run.font.size = Pt(11)

    # Set font size for non-header rows
    for i in range(1, len(rows)):
        for j in range(num_cols):
            cell = table.rows[i].cells[j]
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10)

    doc.add_paragraph()  # spacing after table


def add_code_block(doc, code_lines):
    """Add a code block with Consolas font and gray background."""
    code_text = "\n".join(code_lines)
    p = doc.add_paragraph()
    # Set paragraph background via shading
    pPr = p._p.get_or_add_pPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="F2F2F2"/>')
    pPr.append(shd)

    # Add spacing
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)

    run = p.add_run(code_text)
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)


def add_image(doc, img_path, alt_text):
    """Add an image to the document with caption."""
    if not img_path.exists():
        p = doc.add_paragraph(f"[图片未找到: {img_path}]")
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(img_path), width=Cm(15))

    # Caption below image
    if alt_text:
        cap = doc.add_paragraph(alt_text)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].font.size = Pt(9)
        cap.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        cap.runs[0].italic = True


def add_blockquote(doc, text):
    """Add a blockquote as italic gray paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    apply_inline_formatting(p, text)
    for run in p.runs:
        run.italic = True
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def add_list_item(doc, text):
    """Add a list item paragraph with bullet."""
    p = doc.add_paragraph(style='List Bullet')
    p.clear()
    apply_inline_formatting(p, text)


def convert_md_to_docx(md_path, output_path):
    """Main conversion function."""
    with open(md_path, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    # Set East Asian font
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = parse_xml(f'<w:rFonts {nsdecls("w")}/>') 
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), '微软雅黑')

    # Configure heading styles
    for level in range(1, 4):
        heading_style = doc.styles[f'Heading {level}']
        heading_style.font.name = 'Arial'
        heading_style.font.bold = True
        rPr = heading_style.element.get_or_add_rPr()
        rF = rPr.find(qn('w:rFonts'))
        if rF is None:
            rF = parse_xml(f'<w:rFonts {nsdecls("w")}/>') 
            rPr.append(rF)
        rF.set(qn('w:eastAsia'), '微软雅黑')
        if level == 1:
            heading_style.font.size = Pt(22)
        elif level == 2:
            heading_style.font.size = Pt(16)
        else:
            heading_style.font.size = Pt(13)

    # Page margins (Word default)
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.18)
        section.right_margin = Cm(3.18)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')

        # Skip empty lines
        if not line.strip():
            i += 1
            continue

        # Horizontal rule
        if re.match(r'^---+\s*$', line.strip()):
            add_horizontal_line(doc)
            i += 1
            continue

        # Headings
        heading_match = re.match(r'^(#{1,3})\s+(.*)', line)
        if heading_match:
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            doc.add_heading(text, level=level)
            i += 1
            continue

        # Code block
        if line.strip().startswith('```'):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            add_code_block(doc, code_lines)
            i += 1  # skip closing ```
            continue

        # Image
        img_match = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)', line.strip())
        if img_match:
            alt_text = img_match.group(1)
            img_rel = img_match.group(2)
            img_path = SCREENSHOTS_DIR / img_rel
            add_image(doc, img_path, alt_text)
            i += 1
            continue

        # Blockquote
        if line.strip().startswith('>'):
            text = re.sub(r'^>\s*', '', line.strip())
            add_blockquote(doc, text)
            i += 1
            continue

        # Table
        if '|' in line and i + 1 < len(lines) and re.match(r'^\|[\s\-:]+\|', lines[i + 1].strip()):
            table_lines = []
            while i < len(lines) and '|' in lines[i]:
                table_lines.append(lines[i].rstrip('\n'))
                i += 1
            rows = parse_table(table_lines)
            add_table_to_doc(doc, rows)
            continue

        # List item
        list_match = re.match(r'^[-*]\s+(.*)', line.strip())
        if list_match:
            text = list_match.group(1)
            add_list_item(doc, text)
            i += 1
            continue

        # Ordered list item
        ol_match = re.match(r'^(\d+)\.\s+(.*)', line.strip())
        if ol_match:
            text = ol_match.group(2)
            add_list_item(doc, text)
            i += 1
            continue

        # Normal paragraph - collect consecutive non-empty lines
        para_lines = [line.strip()]
        i += 1
        while i < len(lines):
            next_line = lines[i].rstrip('\n')
            if not next_line.strip():
                break
            if re.match(r'^(#{1,3}\s|```|---+|>|\|.*\||!\[|[-*]\s|\d+\.\s)', next_line.strip()):
                break
            para_lines.append(next_line.strip())
            i += 1

        full_text = " ".join(para_lines)
        p = doc.add_paragraph()
        apply_inline_formatting(p, full_text)

    doc.save(str(output_path))
    print(f"✅ Word document saved to: {output_path}")
    print(f"   File size: {output_path.stat().st_size / 1024:.1f} KB")


if __name__ == "__main__":
    convert_md_to_docx(MD_PATH, OUTPUT_PATH)
